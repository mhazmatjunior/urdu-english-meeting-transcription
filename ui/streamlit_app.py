"""
Streamlit UI — National Language Processing
FYP — M. Hassan Azmat & Wajeeha Islam & Kinza Aslam

Run:
    conda activate fyp
    streamlit run ui/streamlit_app.py
"""

import sys
import json
import threading
import time
from pathlib import Path
from io import BytesIO
import requests

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, str(Path(__file__).parent.parent))

st.set_page_config(
    page_title="National Language Processing",
    page_icon="🗣️",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Serif+Display:ital@0;1&family=DM+Sans:wght@300;400;500;600&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
#MainMenu, footer, header   { visibility: hidden; }
.block-container             { padding-top: 2rem; padding-bottom: 2rem; }

.hero { text-align:center; padding:2.5rem 1rem 1.5rem; border-bottom:1px solid #e8e4dc; margin-bottom:2rem; }
.hero-title { font-family:'DM Serif Display',serif; font-size:2.8rem; color:#1a1714; letter-spacing:-0.02em; margin:0 0 0.4rem; }
.hero-title span { color:#c4622d; font-style:italic; }
.hero-sub { font-size:1rem; color:#7a7168; font-weight:300; margin:0; }

.seg-card { background:#faf9f7; border:1px solid #e8e4dc; border-left:4px solid #c4622d; border-radius:8px; padding:1rem 1.2rem; margin-bottom:0.9rem; }
.seg-card.spk2 { border-left-color:#2d7fc4; }
.seg-card.spk3 { border-left-color:#2dc47a; }
.seg-card.spk4 { border-left-color:#c4b02d; }
.seg-card.spk5 { border-left-color:#9b2dc4; }
.seg-meta { font-size:0.75rem; color:#9a9188; font-weight:600; text-transform:uppercase; letter-spacing:0.08em; margin-bottom:0.5rem; }
.seg-urdu { font-size:1.1rem; color:#1a1714; line-height:1.8; direction:rtl; text-align:right; margin-bottom:0.4rem; }
.seg-english { font-size:0.9rem; color:#5a5450; font-style:italic; line-height:1.6; border-top:1px solid #e8e4dc; padding-top:0.4rem; margin-top:0.2rem; }

.stats-bar { display:flex; gap:2rem; padding:1rem 1.5rem; background:#1a1714; border-radius:10px; margin-bottom:1.5rem; }
.stat-item { text-align:center; }
.stat-val { font-family:'DM Serif Display',serif; font-size:1.8rem; color:#f0ebe3; display:block; }
.stat-label { font-size:0.72rem; color:#7a7168; text-transform:uppercase; letter-spacing:0.1em; }

.stButton > button { background:#c4622d !important; color:white !important; border:none !important; border-radius:8px !important; font-weight:500 !important; padding:0.6rem 1.6rem !important; }
.stButton > button:hover { background:#a84f22 !important; }

.upload-box { background:#faf9f7; border:1px solid #e8e4dc; border-radius:8px; padding:1rem 1.4rem; margin-bottom:1rem; }
.upload-box h4 { margin:0 0 0.5rem; color:#1a1714; font-size:0.95rem; }
.section-heading { font-family:'DM Serif Display',serif; font-size:1.6rem; color:#1a1714; margin-bottom:1rem; padding-bottom:0.4rem; border-bottom:1px solid #e8e4dc; }

[data-testid="column"]:nth-child(7) .stButton > button { background:#c0392b !important; }
[data-testid="column"]:nth-child(7) .stButton > button:hover { background:#a93226 !important; }
</style>
""", unsafe_allow_html=True)

# ════════════════════════════════════════════════════════════
# CORE FUNCTIONS
# ════════════════════════════════════════════════════════════

def merge_consecutive_speakers(segments: list) -> list:
    if not segments:
        return []
    merged  = []
    current = dict(segments[0])
    for seg in segments[1:]:
        if seg.get("speaker") == current.get("speaker"):
            current["end"] = seg["end"]
            urdu_a = current.get("corrected_text") or current.get("text", "")
            urdu_b = seg.get("corrected_text") or seg.get("text", "")
            current["corrected_text"] = (urdu_a + " " + urdu_b).strip()
            eng_a  = current.get("translation_en", "")
            eng_b  = seg.get("translation_en", "")
            current["translation_en"] = (eng_a + " " + eng_b).strip()
        else:
            merged.append(current)
            current = dict(seg)
    merged.append(current)
    return merged

def get_display_name(speaker: str) -> str:
    return st.session_state.speaker_names.get(speaker, speaker)

SPEAKER_CLASSES = ["", "spk2", "spk3", "spk4", "spk5"]

def speaker_class(label: str) -> str:
    try:
        n = int(label.split("_")[-1]) - 1
        return SPEAKER_CLASSES[n % len(SPEAKER_CLASSES)]
    except Exception:
        return ""

def format_time(s: float) -> str:
    return f"{int(s//60)}:{int(s%60):02d}"

def render_segments(segments: list, show_translation: bool = True):
    for block in merge_consecutive_speakers(segments):
        orig_speaker = block.get("speaker", "SPEAKER_?")
        speaker      = get_display_name(orig_speaker)
        css          = speaker_class(orig_speaker)
        urdu         = block.get("corrected_text") or block.get("text", "")
        english      = block.get("translation_en", "")
        tr_html      = f'<div class="seg-english">🇬🇧 {english}</div>' if show_translation and english else ""
        st.markdown(f"""
        <div class="seg-card {css}">
            <div class="seg-meta">{speaker} &nbsp;·&nbsp; {format_time(block.get("start",0))} – {format_time(block.get("end",0))}</div>
            <div class="seg-urdu">{urdu}</div>
            {tr_html}
        </div>""", unsafe_allow_html=True)

def render_stats(segments: list):
    real_spk = set(
        s.get("speaker", "?") for s in segments
        if s.get("speaker", "?") not in ("SPEAKER_UNKNOWN", "SPEAKER_?", "?")
    )
    turns    = len(merge_consecutive_speakers(segments))
    duration = max((s.get("end", 0) for s in segments), default=0)
    with_tr  = sum(1 for s in segments if s.get("translation_en", "").strip())

    st.caption("Speakers: " + ", ".join(get_display_name(spk) for spk in sorted(real_spk)))
    st.markdown(f"""
    <div class="stats-bar">
        <div class="stat-item"><span class="stat-val">{len(segments)}</span><span class="stat-label">Segments</span></div>
        <div class="stat-item"><span class="stat-val">{turns}</span><span class="stat-label">Speaker Turns</span></div>
        <div class="stat-item"><span class="stat-val">{len(real_spk)}</span><span class="stat-label">Speakers</span></div>
        <div class="stat-item"><span class="stat-val">{format_time(duration)}</span><span class="stat-label">Duration</span></div>
        <div class="stat-item"><span class="stat-val">{with_tr}</span><span class="stat-label">Translated</span></div>
    </div>""", unsafe_allow_html=True)

def build_transcript_txt(segments: list) -> str:
    lines = [
        "=" * 70,
        "URDU-ENGLISH MEETING TRANSCRIPTION",
        "National Language Processing — FYP (Wajeeha Islam, Kinza Aslam)",
        "=" * 70, "",
    ]
    for i, block in enumerate(merge_consecutive_speakers(segments), 1):
        lines.append(
            f"[{i:03d}] {get_display_name(block.get('speaker', '?'))}  "
            f"({format_time(block.get('start', 0))} – {format_time(block.get('end', 0))})"
        )
        lines.append(f"  Urdu:    {block.get('corrected_text') or block.get('text', '')}")
        lines.append(f"  English: {block.get('translation_en', '')}")
        lines.append("")
    return "\n".join(lines)

def build_word_document(segments, meeting_title, meeting_date, attendees=""):
    doc = Document()

    title = doc.add_heading(meeting_title or "Meeting Transcription", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"Date: {meeting_date}")
    if attendees:
        meta_run.add_text(f"   |   Attendees: {attendees}")
    meta_run.italic = True

    doc.add_heading("Agenda", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    agenda_para = doc.add_paragraph(meeting_title or "—")
    agenda_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_heading("Transcription", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER

    for block in merge_consecutive_speakers(segments):
        speaker = get_display_name(block.get("speaker", "Unknown"))
        urdu    = block.get("corrected_text") or block.get("text", "")
        english = block.get("translation_en", "")

        speaker_para = doc.add_paragraph()
        speaker_run = speaker_para.add_run(
            f"{speaker}  ({format_time(block.get('start', 0))} – {format_time(block.get('end', 0))})"
        )
        speaker_run.bold = True
        speaker_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        urdu_para = doc.add_paragraph(urdu)
        urdu_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        if english:
            eng_para = doc.add_paragraph()
            eng_run = eng_para.add_run(english)
            eng_run.italic = True
            eng_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def run_tts(segments: list):
    tts_dir = Path("data/output/audio_segments")
    tts_dir.mkdir(parents=True, exist_ok=True)

    payload = {"segments": segments}
    resp = requests.post("http://127.0.0.1:8000/synthesize", json=payload, timeout=600)
    resp.raise_for_status()

    audio_bytes = resp.content
    if not audio_bytes:
        return None

    path = "data/output/audio_segments/full_audio.wav"
    with open(path, "wb") as f:
        f.write(audio_bytes)
    return path

# ════════════════════════════════════════════════════════════
# SESSION STATE
# ════════════════════════════════════════════════════════════

for key, val in [
    ("segments",      None),
    ("processed",     False),
    ("audio_path",    None),
    ("show_tr",       True),
    ("last_file_id",  None),
    ("speaker_names", {}),
    ("show_rename",   False),
    ("meeting_title", ""),
    ("meeting_date",  None),
    ("attendees",     ""),
]:
    if key not in st.session_state:
        st.session_state[key] = val

# ════════════════════════════════════════════════════════════
# HERO
# ════════════════════════════════════════════════════════════

st.markdown("""
<div class="hero">
    <h1 class="hero-title">National <span>Language Processing</span></h1>
    <p class="hero-sub">Urdu · English · Speaker Diarization · Meeting Transcription</p>
</div>""", unsafe_allow_html=True)

# ════════════════════════════════════════════════════════════
# UPLOAD
# ════════════════════════════════════════════════════════════

if not st.session_state.processed:
    st.markdown("""
    <div class="upload-box">
        <h4>📂 Upload Meeting Recording</h4>
        Upload the audio or video file of your meeting.
        The system will automatically transcribe the Urdu speech, identify speakers, correct the text, and provide an English translation.
    </div>""", unsafe_allow_html=True)

    uploaded_audio = st.file_uploader(
        "Upload audio or video",
        type=["wav", "mp3", "mp4", "m4a"],
        key="audio_upload",
    )

    if uploaded_audio is not None:
        if st.session_state.last_file_id != uploaded_audio.file_id:
            st.session_state.last_file_id = uploaded_audio.file_id

            result_holder = {}

            def call_api():
                try:
                    files = {"audio": (uploaded_audio.name, uploaded_audio.getvalue())}
                    resp = requests.post(
                        "http://127.0.0.1:8000/transcribe",
                        files=files,
                        timeout=600,  # 10 minute — Modal cold-start ke liye kaafi waqt
                    )
                    resp.raise_for_status()
                    result_holder["data"] = resp.json()
                except Exception as e:
                    result_holder["error"] = str(e)

            thread = threading.Thread(target=call_api)
            thread.start()

            status_messages = [
                "🎧 Listening...",
                "🗣️ Loading Whisper...",
                "👥 Speakers Diarization...",
                "✍️ Text transcription and correction...",
                "🌐 English Translation...",
                "⏳  Almost there...",
                "✅ Finalizing results...",
                "📦 Preparing output files...",
                "Just a few more seconds...",
            ]
            msg_box = st.empty()
            progress_bar = st.progress(0)
            start = time.time()
            estimated_seconds = 90
            message_interval = 8  # message har 8 second baad badlega

            i = 0
            last_message_time = -message_interval  # taake pehla message turant dikhe
            while thread.is_alive():
                elapsed = time.time() - start

                if elapsed - last_message_time >= message_interval:
                    msg_box.info(status_messages[i % len(status_messages)])
                    i += 1
                    last_message_time = elapsed

                progress_bar.progress(min(elapsed / estimated_seconds, 0.95))
                time.sleep(1)

            thread.join()
            progress_bar.progress(1.0)
            msg_box.empty()

            if "error" in result_holder:
                st.error(f"Transcription fail ho gayi: {result_holder['error']}")
                st.stop()

            result = result_holder["data"]
            st.session_state.segments = result["segments"]
            st.session_state.processed = True

            Path("data/output").mkdir(parents=True, exist_ok=True)
            with open("data/output/corrected_transcription.json", "w", encoding="utf-8") as f:
                json.dump({"segments": st.session_state.segments}, f, ensure_ascii=False, indent=2)

            st.rerun()

# ════════════════════════════════════════════════════════════
# RESULTS
# ════════════════════════════════════════════════════════════

if st.session_state.processed and st.session_state.segments:
    segments = st.session_state.segments

    render_stats(segments)

    with st.expander("📝 Meeting Details (for Word Document)", expanded=not st.session_state.meeting_title):
        st.text_input("Meeting Title / Agenda", placeholder="e.g. Weekly FYP Progress Sync", key="meeting_title")
        st.date_input("Meeting Date", key="meeting_date")
        st.text_input("Attendees (optional)", placeholder="e.g. Wajeeha, Kinza", key="attendees")

    c1, c2, c3, c4, c5, c6, c7 = st.columns(7)

    with c1:
        st.session_state.show_tr = st.toggle("🌐 Translation", value=st.session_state.show_tr)

    with c2:
        if st.button("✏️ Rename", use_container_width=True):
            st.session_state.show_rename = not st.session_state.show_rename

    with c3:
        tts_btn = st.button("🔊 Generate Audio", use_container_width=True)

    with c4:
        st.download_button(
            "⬇ results.json",
            data=json.dumps({"segments": segments}, ensure_ascii=False, indent=2),
            file_name="results.json",
            mime="application/json",
            use_container_width=True,
        )

    with c5:
        st.download_button(
            "⬇ transcript.txt",
            data=build_transcript_txt(segments),
            file_name="transcript.txt",
            mime="text/plain",
            use_container_width=True,
        )

    with c6:
        docx_buffer = build_word_document(
            segments,
            st.session_state.meeting_title,
            st.session_state.meeting_date,
            st.session_state.attendees,
        )
        st.download_button(
            "⬇ Word Doc",
            data=docx_buffer,
            file_name="meeting_transcript.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    with c7:
        if st.button("🔄 New Transcription", use_container_width=True):
            st.session_state.segments      = None
            st.session_state.processed     = False
            st.session_state.audio_path    = None
            st.session_state.last_file_id  = None
            st.session_state.speaker_names = {}
            st.rerun()

    # Rename panel
    if st.session_state.show_rename:
        st.markdown("---")
        all_spk = sorted(set(
            s.get("speaker", "?") for s in segments
            if s.get("speaker", "?") not in ("SPEAKER_UNKNOWN", "SPEAKER_?", "?")
        ))
        cols = st.columns(len(all_spk)) if len(all_spk) <= 4 else st.columns(4)
        for i, spk in enumerate(all_spk):
            with cols[i % 4]:
                current_name = st.session_state.speaker_names.get(spk, "")
                new_name = st.text_input(f"{spk}", value=current_name,
                                         placeholder="e.g. Ahmed", key=f"rename_{spk}")
                if new_name.strip():
                    st.session_state.speaker_names[spk] = new_name.strip()
                elif spk in st.session_state.speaker_names:
                    del st.session_state.speaker_names[spk]
        st.markdown("---")

    # TTS
    if tts_btn:
        with st.spinner("Generating English audio..."):
            st.session_state.audio_path = run_tts(segments)

    if st.session_state.audio_path and Path(st.session_state.audio_path).exists():
        st.markdown("#### English Audio")
        st.audio(st.session_state.audio_path)
        with open(st.session_state.audio_path, "rb") as f:
            st.download_button("⬇ Download full_audio.wav", data=f,
                               file_name="full_audio.wav", mime="audio/wav")

    # Transcription feed
    st.markdown('<div class="section-heading">Transcription</div>', unsafe_allow_html=True)

    all_speakers    = sorted(set(
        s.get("speaker", "?") for s in segments
        if s.get("speaker", "?") not in ("SPEAKER_UNKNOWN", "SPEAKER_?", "?")
    ))
    display_options = {get_display_name(spk): spk for spk in all_speakers}
    selected_display = st.multiselect(
        "Filter by speaker",
        options=list(display_options.keys()),
        default=list(display_options.keys()),
    )
    selected = [display_options[d] for d in selected_display]
    filtered = [s for s in segments if s.get("speaker", "?") in selected]

    render_segments(filtered, show_translation=st.session_state.show_tr)

    turns = len(merge_consecutive_speakers(filtered))
    st.caption(f"Showing {turns} speaker turns ({len(filtered)} segments) of {len(segments)} total")